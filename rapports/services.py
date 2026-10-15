"""
Services de génération des rapports PDF et Excel.

PDF : WeasyPrint, généré en mémoire (BytesIO).
Excel : openpyxl, généré en mémoire (BytesIO).
"""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.db.models import Q, Sum
from django.template.loader import render_to_string
from django.utils import timezone
from weasyprint import HTML
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from agents.models import Agent
from core.models import Parametre
from audit.models import JournalAudit
from decomptes.models import Decompte
from gardes.models import Garde
from paiements.models import OrdrePaiement, Paiement


# ===================================================================
# Helpers
# ===================================================================

def _date_edition():
    return timezone.localtime(timezone.now())


def _reference_audit(entite, id_entite, action=JournalAudit.ACTION_VALIDATION):
    """Retourne la dernière entrée d'audit correspondant à (entite, id)."""
    return (
        JournalAudit.objects
        .filter(
            entite_concernee=entite,
            id_entite_concernee=id_entite,
            action=action,
        )
        .order_by('-date_action')
        .first()
    )


# ===================================================================
# PDF génériques
# ===================================================================

def _pdf_from_template(template_name, context, base_url=None):
    """Rend un template et retourne les bytes PDF."""
    html = render_to_string(template_name, context)
    pdf_bytes = HTML(string=html, base_url=base_url).write_pdf()
    return pdf_bytes


# ===================================================================
# Rapport 1 — État de décompte
# ===================================================================

def generer_etat_decompte_pdf(decompte):
    """
    Génère le PDF de l'état de décompte individuel d'un agent,
    au format officiel (colonnes début/fin/matricule/partiels/montant/total).
    """
    lignes = list(
        decompte.lignes
        .select_related('garde', 'tarif', 'garde__agent')
        .order_by('garde__date_heure_debut')
    )

    total_partiels = sum(l.nombre_partiels_snapshot for l in lignes)
    total_montant = sum(l.montant_partiel_snapshot for l in lignes)

    audit_entry = _reference_audit(
        'decomptes.Decompte', decompte.pk, JournalAudit.ACTION_VALIDATION
    )

    context = {
        'decompte': decompte,
        'lignes': lignes,
        'total_partiels': total_partiels,
        'total_montant': total_montant,
        'audit_entry': audit_entry,
        'date_edition': _date_edition(),
    }

    pdf_bytes = _pdf_from_template(
        'rapports/pdf/etat_decompte.html', context
    )
    return BytesIO(pdf_bytes)


def collecter_decomptes_mensuels(annee, mois):
    """
    Retourne la liste des groupes { agent, lignes, totaux } pour un mois donné,
    uniquement pour les décomptes VALIDÉS.
    """
    decomptes = (
        Decompte.objects
        .filter(annee=annee, mois=mois, statut=Decompte.STATUT_VALIDE)
        .select_related('agent')
        .order_by('agent__nom', 'agent__prenoms')
    )

    groupes = []
    grand_total_partiels = 0
    grand_total_montant = Decimal('0.00')
    grand_total = Decimal('0.00')

    for d in decomptes:
        lignes = list(
            d.lignes
            .select_related('garde')
            .order_by('garde__date_heure_debut')
        )
        total_partiels = sum(l.nombre_partiels_snapshot for l in lignes)
        total_montant = sum(l.montant_partiel_snapshot for l in lignes)
        total_general = d.montant_total

        groupes.append({
            'agent': d.agent,
            'decompte': d,
            'lignes': lignes,
            'total_partiels': total_partiels,
            'total_montant': total_montant,
            'total_general': total_general,
        })

        grand_total_partiels += total_partiels
        grand_total_montant += total_montant
        grand_total += total_general

    return {
        'groupes': groupes,
        'grand_total_partiels': grand_total_partiels,
        'grand_total_montant': grand_total_montant,
        'grand_total': grand_total,
    }


def generer_decomptes_mensuels_pdf(annee, mois):
    """
    Génère le PDF récapitulatif mensuel (tous les agents d'un mois),
    au format officiel avec sous-totaux par agent et total général.
    """
    data = collecter_decomptes_mensuels(annee, mois)
    context = {
        'annee': annee,
        'mois': mois,
        **data,
        'date_edition': _date_edition(),
    }
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/decomptes_mensuels.html', context
    )
    return BytesIO(pdf_bytes)


# ===================================================================
# Rapport 2 — Liste de paiement
# ===================================================================

def _param_op(cle, defaut=""):
    """Lit un paramètre OP depuis core.Parametre. Retourne `defaut` si absent."""
    try:
        return Parametre.objects.get(cle=cle).valeur
    except Parametre.DoesNotExist:
        return defaut


def _nom_mois_fr(numero):
    """1 → 'Janvier'."""
    from rapports.templatetags.format_tags import MOIS_FR
    return MOIS_FR.get(int(numero), '')


def generer_liste_paiement_pdf(op):
    """
    Génère le PDF de l'Ordre de Paiement au format officiel.

    Lit les valeurs stables depuis core.Parametre.
    Gère le fallback titulaire/adresse/matricule depuis l'agent bénéficiaire.
    """
    agent = op.agent_beneficiaire

    # Fallback titulaire
    titulaire_nom = op.titulaire_nom or f"{agent.nom} {agent.prenoms}"
    titulaire_numero = op.titulaire_numero or agent.matricule
    cin_titulaire = agent.cin or ""
    adresse_titulaire = agent.adresse or ""

    # Période couverte
    periode_debut = op.periodes.filter(reservation_active=True).order_by(
        'decompte__annee', 'decompte__mois'
    ).first()
    periode_fin = op.periodes.filter(reservation_active=True).order_by(
        '-decompte__annee', '-decompte__mois'
    ).first()

    if periode_debut and periode_fin:
        mois_debut = _nom_mois_fr(periode_debut.decompte.mois)
        mois_fin = _nom_mois_fr(periode_fin.decompte.mois)
        annee_debut = periode_debut.decompte.annee
        annee_fin = periode_fin.decompte.annee
        if annee_debut == annee_fin:
            periode_texte = f"{mois_debut} à {mois_fin} {annee_fin}"
        else:
            periode_texte = f"{mois_debut} {annee_debut} à {mois_fin} {annee_fin}"
    else:
        periode_texte = "—"

    # QR code avec les informations clés de l'OP
    qr_data = (
        f"OP: {op.numero_op}\n"
        f"Agent: {agent.matricule} - {agent.nom} {agent.prenoms}\n"
        f"Montant: {op.montant_total} Ar\n"
        f"Émis le: {op.date_creation.strftime('%d/%m/%Y')}"
    )
    qr_data_url = _generer_qr_base64(qr_data, taille_px=4)

    context = {
        'op': op,
        'agent': agent,
        'qr_code': qr_data_url,
        'date_edition': _date_edition(),

        # Titulaire (fallback)
        'titulaire_nom': titulaire_nom,
        'titulaire_numero': titulaire_numero,
        'cin_titulaire': cin_titulaire,
        'adresse_titulaire': adresse_titulaire,

        # Période
        'periode_texte': periode_texte,

        # Paramètres officiels (16 valeurs)
        'p_ministere': _param_op('op.ministere', '21 MEF'),
        'p_ordonnateur': _param_op('op.ordonnateur', '00-221-A-000000'),
        'p_commune_beneficiaire': _param_op('op.commune_beneficiaire', '00000'),
        'p_code_cf': _param_op('op.code_cf', '0101 220'),
        'p_tp_assignataire': _param_op('op.tp_assignataire', '10101 100'),
        'p_programme': _param_op('op.programme', '117'),
        'p_activite_projet': _param_op('op.activite_projet', '117-1-1-1-01'),
        'p_imputation_admin': _param_op('op.imputation_admin', '00-21-7-223-00000'),
        'p_compte': _param_op('op.compte', '6241 Indemnités de mission intérieure'),
        'p_republique': _param_op('op.republique', 'REPUBLIQUE DE MADAGASCAR'),
        'p_code_mission': _param_op('op.code_mission', '220'),
        'p_budget': _param_op('op.budget', '01'),
        'p_categorie': _param_op('op.categorie', '3'),
        'p_sect_conv': _param_op('op.sect_conv', '000 RESSOURCES PROPRES'),
        'p_financement': _param_op('op.financement', '110-001-001-A'),
        'p_pieces_justificatives': _param_op(
            'op.pieces_justificatives',
            'Ordre de mission — État de décompte'
        ),
    }

    pdf_bytes = _pdf_from_template(
        'rapports/pdf/liste_paiement.html', context
    )
    return BytesIO(pdf_bytes)


# ===================================================================
# Rapport 3 — Historique agent
# ===================================================================

def _bornes_periode(date_debut, date_fin):
    """
    Retourne (debut, fin) en `date`.
    Si non spécifiée : 12 derniers mois.
    """
    if date_fin is None:
        date_fin = date.today()
    if date_debut is None:
        date_debut = date_fin - timedelta(days=365)
    if date_debut > date_fin:
        raise ValueError(
            "La date de début doit être antérieure ou égale à la date de fin."
        )
    return date_debut, date_fin


def collecter_donnees_historique_agent(agent, date_debut, date_fin):
    """
    Retourne un dictionnaire contenant les données consolidées
    pour un agent sur la période.

    Règles :
    - Gardes : 30 dernières, début dans la période.
    - Décomptes : 12 derniers, dont l'année/mois tombe dans la période.
    - OP : 12 derniers, date_creation dans la période.
    - Paiements : tous les paiements des OP concernés.
    """
    debut, fin = _bornes_periode(date_debut, date_fin)

    gardes = (
        Garde.objects
        .filter(
            agent=agent,
            date_heure_debut__date__gte=debut,
            date_heure_debut__date__lte=fin,
        )
        .select_related('mission', 'tarif')
        .order_by('-date_heure_debut')[:30]
    )

    decomptes = (
        Decompte.objects
        .filter(agent=agent)
        .filter(
            Q(annee__gt=debut.year)
            | Q(annee=debut.year, mois__gte=debut.month)
        )
        .filter(
            Q(annee__lt=fin.year)
            | Q(annee=fin.year, mois__lte=fin.month)
        )
        .order_by('-annee', '-mois')[:12]
    )

    ordres = (
        OrdrePaiement.objects
        .filter(
            agent_beneficiaire=agent,
            date_creation__date__gte=debut,
            date_creation__date__lte=fin,
        )
        .order_by('-date_creation')[:12]
    )

    paiements = (
        Paiement.objects
        .filter(ordre_paiement__in=ordres)
        .select_related('ordre_paiement')
        .order_by('-date_paiement')
    )

    return {
        'agent': agent,
        'debut': debut,
        'fin': fin,
        'gardes': gardes,
        'decomptes': decomptes,
        'ordres': ordres,
        'paiements': paiements,
    }


def generer_historique_agent_pdf(agent, date_debut=None, date_fin=None):
    """Génère le PDF de l'historique d'un agent."""
    data = collecter_donnees_historique_agent(agent, date_debut, date_fin)
    data['date_edition'] = _date_edition()
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/historique_agent.html', data
    )
    return BytesIO(pdf_bytes)


# ===================================================================
# Rapport 3 — Historique agent (Excel)
# ===================================================================

_ENTETE_STYLE = Font(bold=True, color="FFFFFF")
_ENTETE_FILL = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
_BORDER = Border(
    left=Side(style='thin', color='CCCCCC'),
    right=Side(style='thin', color='CCCCCC'),
    top=Side(style='thin', color='CCCCCC'),
    bottom=Side(style='thin', color='CCCCCC'),
)


def _ecrire_entete(ws, titres, row=1):
    for col_idx, titre in enumerate(titres, start=1):
        cell = ws.cell(row=row, column=col_idx, value=titre)
        cell.font = _ENTETE_STYLE
        cell.fill = _ENTETE_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER


def _ajuster_colonnes(ws):
    for col in ws.columns:
        max_len = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 3, 40)


def generer_historique_agent_excel(agent, date_debut=None, date_fin=None):
    """Génère le classeur Excel (.xlsx) de l'historique d'un agent."""
    data = collecter_donnees_historique_agent(agent, date_debut, date_fin)

    wb = Workbook()

    # ------- Feuille 1 : Informations générales -------
    ws_info = wb.active
    ws_info.title = "Agent"
    ws_info["A1"] = "Fiche agent"
    ws_info["A1"].font = Font(bold=True, size=14)
    ws_info["A3"] = "Matricule"
    ws_info["B3"] = agent.matricule
    ws_info["A4"] = "Nom"
    ws_info["B4"] = agent.nom
    ws_info["A5"] = "Prénoms"
    ws_info["B5"] = agent.prenoms
    ws_info["A6"] = "Grade"
    ws_info["B6"] = agent.grade or "—"
    ws_info["A7"] = "Statut"
    ws_info["B7"] = agent.get_statut_agent_display()
    ws_info["A8"] = "Période analysée"
    ws_info["B8"] = f"{data['debut'].isoformat()} → {data['fin'].isoformat()}"
    ws_info["A9"] = "Date d'édition"
    ws_info["B9"] = _date_edition().strftime("%d/%m/%Y %H:%M")
    _ajuster_colonnes(ws_info)

    # ------- Feuille 2 : Gardes -------
    ws_g = wb.create_sheet("Gardes")
    _ecrire_entete(ws_g, [
        "Garde #", "Début", "Fin", "Durée (min)",
        "Partiels", "Conforme", "Tarif (Ar)", "Statut",
    ])
    for i, g in enumerate(data['gardes'], start=2):
        ws_g.cell(row=i, column=1, value=g.pk)
        ws_g.cell(row=i, column=2, value=g.date_heure_debut.replace(tzinfo=None))
        ws_g.cell(row=i, column=3, value=g.date_heure_fin.replace(tzinfo=None))
        ws_g.cell(row=i, column=4, value=g.duree_minutes)
        ws_g.cell(row=i, column=5, value=g.nombre_partiels)
        ws_g.cell(row=i, column=6, value="Oui" if g.conforme else "Non")
        ws_g.cell(row=i, column=7, value=float(g.tarif.montant_partiel) if g.tarif else None)
        ws_g.cell(row=i, column=8, value=g.get_statut_display())
        for col in range(1, 9):
            ws_g.cell(row=i, column=col).border = _BORDER
        ws_g.cell(row=i, column=2).number_format = 'DD/MM/YYYY HH:MM'
        ws_g.cell(row=i, column=3).number_format = 'DD/MM/YYYY HH:MM'
        ws_g.cell(row=i, column=7).number_format = '#,##0.00 "Ar"'
    _ajuster_colonnes(ws_g)

    # ------- Feuille 3 : Décomptes -------
    ws_d = wb.create_sheet("Decomptes")
    _ecrire_entete(ws_d, [
        "Décompte #", "Période", "Gardes retenues",
        "Non conformes exclues", "Montant total (Ar)", "Statut",
    ])
    for i, d in enumerate(data['decomptes'], start=2):
        ws_d.cell(row=i, column=1, value=d.pk)
        ws_d.cell(row=i, column=2, value=f"{d.mois:02d}/{d.annee}")
        ws_d.cell(row=i, column=3, value=d.nombre_gardes_retunes)
        ws_d.cell(row=i, column=4, value=d.nombre_gardes_non_conformes_exclues)
        ws_d.cell(row=i, column=5, value=float(d.montant_total))
        ws_d.cell(row=i, column=6, value=d.get_statut_display())
        for col in range(1, 7):
            ws_d.cell(row=i, column=col).border = _BORDER
        ws_d.cell(row=i, column=5).number_format = '#,##0.00 "Ar"'
    _ajuster_colonnes(ws_d)

    # ------- Feuille 4 : OP et Paiements -------
    ws_op = wb.create_sheet("OP et Paiements")
    _ecrire_entete(ws_op, [
        "N° OP", "Date OP", "Montant total (Ar)", "Net payé (Ar)",
        "Statut OP", "Clôturé", "Paiement date", "Paiement montant (Ar)",
        "Mode paiement", "Statut paiement",
    ])
    row = 2
    for op in data['ordres']:
        paiements_op = list(op.paiements.all())
        if not paiements_op:
            paiements_op = [None]
        for p in paiements_op:
            ws_op.cell(row=row, column=1, value=op.numero_op)
            ws_op.cell(row=row, column=2, value=op.date_creation.replace(tzinfo=None))
            ws_op.cell(row=row, column=3, value=float(op.montant_total))
            ws_op.cell(row=row, column=4, value=float(op.montant_net_paye))
            ws_op.cell(row=row, column=5, value=op.get_statut_display())
            ws_op.cell(row=row, column=6, value="Oui" if op.cloture_administrative else "Non")
            if p:
                ws_op.cell(row=row, column=7, value=p.date_paiement.replace(tzinfo=None))
                ws_op.cell(row=row, column=8, value=float(p.montant))
                ws_op.cell(row=row, column=9, value=p.get_mode_paiement_display())
                ws_op.cell(row=row, column=10, value=p.get_statut_display())
                ws_op.cell(row=row, column=7).number_format = 'DD/MM/YYYY HH:MM'
                ws_op.cell(row=row, column=8).number_format = '#,##0.00 "Ar"'
            for col in range(1, 11):
                ws_op.cell(row=row, column=col).border = _BORDER
            ws_op.cell(row=row, column=2).number_format = 'DD/MM/YYYY HH:MM'
            ws_op.cell(row=row, column=3).number_format = '#,##0.00 "Ar"'
            ws_op.cell(row=row, column=4).number_format = '#,##0.00 "Ar"'
            row += 1
    _ajuster_colonnes(ws_op)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ===================================================================
# Rapport 4 — État détaillé des gardes
# ===================================================================

def collecter_gardes_filtrees(date_debut=None, date_fin=None,
                              agent=None, mission=None, conforme=None):
    """
    Retourne les gardes filtrées sur la période et critères donnés.
    Les gardes ANNULÉES sont toujours incluses pour traçabilité
    mais identifiées par leur statut.
    """
    qs = (
        Garde.objects
        .select_related('agent', 'mission', 'tarif', 'id_utilisateur_saisie')
        .order_by('-date_heure_debut')
    )
    if date_debut:
        qs = qs.filter(date_heure_debut__date__gte=date_debut)
    if date_fin:
        qs = qs.filter(date_heure_debut__date__lte=date_fin)
    if agent:
        qs = qs.filter(agent=agent)
    if mission:
        qs = qs.filter(mission=mission)
    if conforme is True:
        qs = qs.filter(conforme=True)
    elif conforme is False:
        qs = qs.filter(conforme=False)
    return qs


def generer_etat_gardes_pdf(date_debut=None, date_fin=None,
                            agent=None, mission=None, conforme=None):
    gardes = collecter_gardes_filtrees(
        date_debut, date_fin, agent, mission, conforme
    )
    context = {
        'gardes': gardes,
        'date_debut': date_debut,
        'date_fin': date_fin,
        'agent_filtre': agent,
        'mission_filtre': mission,
        'conforme_filtre': conforme,
        'total': gardes.count(),
        'date_edition': _date_edition(),
    }
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/etat_gardes.html', context
    )
    return BytesIO(pdf_bytes)


def generer_etat_gardes_excel(date_debut=None, date_fin=None,
                              agent=None, mission=None, conforme=None):
    gardes = collecter_gardes_filtrees(
        date_debut, date_fin, agent, mission, conforme
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Gardes"
    _ecrire_entete(ws, [
        "Garde #", "Agent", "Matricule", "Mission",
        "Début", "Fin", "Durée (min)", "Partiels",
        "Conforme", "Statut", "Tarif (Ar)",
    ])
    for i, g in enumerate(gardes, start=2):
        ws.cell(row=i, column=1, value=g.pk)
        ws.cell(row=i, column=2, value=f"{g.agent.nom} {g.agent.prenoms}")
        ws.cell(row=i, column=3, value=g.agent.matricule)
        ws.cell(row=i, column=4, value=g.mission.libelle if g.mission else "—")
        ws.cell(row=i, column=5, value=g.date_heure_debut.replace(tzinfo=None))
        ws.cell(row=i, column=6, value=g.date_heure_fin.replace(tzinfo=None))
        ws.cell(row=i, column=7, value=g.duree_minutes)
        ws.cell(row=i, column=8, value=g.nombre_partiels)
        ws.cell(row=i, column=9, value="Oui" if g.conforme else "Non")
        ws.cell(row=i, column=10, value=g.get_statut_display())
        ws.cell(row=i, column=11, value=float(g.tarif.montant_partiel) if g.tarif else None)
        for col in range(1, 12):
            ws.cell(row=i, column=col).border = _BORDER
        ws.cell(row=i, column=5).number_format = 'DD/MM/YYYY HH:MM'
        ws.cell(row=i, column=6).number_format = 'DD/MM/YYYY HH:MM'
        ws.cell(row=i, column=11).number_format = '#,##0.00 "Ar"'
    _ajuster_colonnes(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ===================================================================
# Rapport 5 — Liste des OP
# ===================================================================

def collecter_ops_filtres(date_debut=None, date_fin=None, statut=None):
    qs = (
        OrdrePaiement.objects
        .select_related('agent_beneficiaire', 'id_utilisateur_createur',
                        'id_utilisateur_validateur')
        .order_by('-annee', '-numero_op')
    )
    if date_debut:
        qs = qs.filter(date_creation__date__gte=date_debut)
    if date_fin:
        qs = qs.filter(date_creation__date__lte=date_fin)
    if statut in dict(OrdrePaiement.STATUTS_CHOICES):
        qs = qs.filter(statut=statut)
    return qs


def generer_liste_op_pdf(date_debut=None, date_fin=None, statut=None):
    ops = collecter_ops_filtres(date_debut, date_fin, statut)
    total_montant = sum((op.montant_total for op in ops), Decimal('0.00'))
    total_paye = sum((op.montant_net_paye for op in ops), Decimal('0.00'))

    context = {
        'ops': ops,
        'date_debut': date_debut,
        'date_fin': date_fin,
        'statut_filtre': statut,
        'total': ops.count(),
        'total_montant': total_montant,
        'total_paye': total_paye,
        'date_edition': _date_edition(),
    }
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/liste_op.html', context
    )
    return BytesIO(pdf_bytes)


def generer_liste_op_excel(date_debut=None, date_fin=None, statut=None):
    ops = collecter_ops_filtres(date_debut, date_fin, statut)
    wb = Workbook()
    ws = wb.active
    ws.title = "Ordres de paiement"
    _ecrire_entete(ws, [
        "N° OP", "Agent", "Matricule", "Année",
        "Montant total (Ar)", "Net payé (Ar)",
        "Statut", "Clôturé", "Créé le", "Créé par",
    ])
    for i, op in enumerate(ops, start=2):
        ws.cell(row=i, column=1, value=op.numero_op)
        ws.cell(row=i, column=2, value=f"{op.agent_beneficiaire.nom} {op.agent_beneficiaire.prenoms}")
        ws.cell(row=i, column=3, value=op.agent_beneficiaire.matricule)
        ws.cell(row=i, column=4, value=op.annee)
        ws.cell(row=i, column=5, value=float(op.montant_total))
        ws.cell(row=i, column=6, value=float(op.montant_net_paye))
        ws.cell(row=i, column=7, value=op.get_statut_display())
        ws.cell(row=i, column=8, value="Oui" if op.cloture_administrative else "Non")
        ws.cell(row=i, column=9, value=op.date_creation.replace(tzinfo=None))
        ws.cell(row=i, column=10, value=op.id_utilisateur_createur.username)
        for col in range(1, 11):
            ws.cell(row=i, column=col).border = _BORDER
        ws.cell(row=i, column=5).number_format = '#,##0.00 "Ar"'
        ws.cell(row=i, column=6).number_format = '#,##0.00 "Ar"'
        ws.cell(row=i, column=9).number_format = 'DD/MM/YYYY HH:MM'
    _ajuster_colonnes(ws)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ===================================================================
# Rapport 6 — Rapport mensuel
# ===================================================================

def collecter_donnees_mensuelles(annee, mois):
    decomptes = (
        Decompte.objects
        .filter(annee=annee, mois=mois, statut=Decompte.STATUT_VALIDE)
        .select_related('agent', 'id_utilisateur_validateur')
        .order_by('agent__nom', 'agent__prenoms')
    )
    total_decomptes = sum((d.montant_total for d in decomptes), Decimal('0.00'))

    # Paiements du mois
    paiements = (
        Paiement.objects
        .filter(
            date_paiement__year=annee,
            date_paiement__month=mois,
            statut=Paiement.STATUT_ENREGISTRE,
        )
        .select_related('ordre_paiement', 'ordre_paiement__agent_beneficiaire')
        .order_by('-date_paiement')
    )
    total_paiements = sum((p.montant for p in paiements), Decimal('0.00'))

    return {
        'annee': annee,
        'mois': mois,
        'decomptes': decomptes,
        'paiements': paiements,
        'total_decomptes': total_decomptes,
        'total_paiements': total_paiements,
        'nb_decomptes': decomptes.count(),
        'nb_paiements': paiements.count(),
    }


def generer_rapport_mensuel_pdf(annee, mois):
    data = collecter_donnees_mensuelles(annee, mois)
    data['date_edition'] = _date_edition()
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/rapport_mensuel.html', data
    )
    return BytesIO(pdf_bytes)


# ===================================================================
# Rapport 7 — Rapport sur période
# ===================================================================

def collecter_donnees_periode(date_debut, date_fin):
    gardes = Garde.objects.filter(
        date_heure_debut__date__gte=date_debut,
        date_heure_debut__date__lte=date_fin,
        statut=Garde.STATUT_VALIDEE,
    )
    decomptes = Decompte.objects.filter(
        date_validation__date__gte=date_debut,
        date_validation__date__lte=date_fin,
        statut=Decompte.STATUT_VALIDE,
    )
    ops = OrdrePaiement.objects.filter(
        date_creation__date__gte=date_debut,
        date_creation__date__lte=date_fin,
    )
    paiements = Paiement.objects.filter(
        date_paiement__date__gte=date_debut,
        date_paiement__date__lte=date_fin,
        statut=Paiement.STATUT_ENREGISTRE,
    )

    return {
        'date_debut': date_debut,
        'date_fin': date_fin,
        'nb_gardes': gardes.count(),
        'nb_decomptes': decomptes.count(),
        'nb_ops': ops.count(),
        'nb_paiements': paiements.count(),
        'total_decomptes': sum((d.montant_total for d in decomptes), Decimal('0.00')),
        'total_ops': sum((op.montant_total for op in ops), Decimal('0.00')),
        'total_paiements': sum((p.montant for p in paiements), Decimal('0.00')),
        'decomptes': decomptes.select_related('agent'),
        'ops': ops.select_related('agent_beneficiaire'),
        'paiements': paiements.select_related('ordre_paiement'),
    }


def generer_rapport_periode_pdf(date_debut, date_fin):
    data = collecter_donnees_periode(date_debut, date_fin)
    data['date_edition'] = _date_edition()
    pdf_bytes = _pdf_from_template(
        'rapports/pdf/rapport_periode.html', data
    )
    return BytesIO(pdf_bytes)


def generer_rapport_periode_excel(date_debut, date_fin):
    data = collecter_donnees_periode(date_debut, date_fin)
    wb = Workbook()

    # Feuille synthèse
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = "Rapport sur période"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A3"] = "Du"
    ws["B3"] = date_debut.isoformat()
    ws["A4"] = "Au"
    ws["B4"] = date_fin.isoformat()
    ws["A6"] = "Nombre de gardes"
    ws["B6"] = data['nb_gardes']
    ws["A7"] = "Nombre de décomptes"
    ws["B7"] = data['nb_decomptes']
    ws["A8"] = "Nombre d'OP"
    ws["B8"] = data['nb_ops']
    ws["A9"] = "Nombre de paiements"
    ws["B9"] = data['nb_paiements']
    ws["A10"] = "Total décomptes (Ar)"
    ws["B10"] = float(data['total_decomptes'])
    ws["A11"] = "Total OP (Ar)"
    ws["B11"] = float(data['total_ops'])
    ws["A12"] = "Total paiements (Ar)"
    ws["B12"] = float(data['total_paiements'])
    _ajuster_colonnes(ws)

    # Feuille décomptes
    ws_d = wb.create_sheet("Décomptes")
    _ecrire_entete(ws_d, ["Agent", "Période", "Montant (Ar)"])
    for i, d in enumerate(data['decomptes'], start=2):
        ws_d.cell(row=i, column=1, value=f"{d.agent.matricule} {d.agent.nom}")
        ws_d.cell(row=i, column=2, value=f"{d.mois:02d}/{d.annee}")
        ws_d.cell(row=i, column=3, value=float(d.montant_total))
        ws_d.cell(row=i, column=3).number_format = '#,##0.00 "Ar"'
        for col in range(1, 4):
            ws_d.cell(row=i, column=col).border = _BORDER
    _ajuster_colonnes(ws_d)

    # Feuille OP
    ws_op = wb.create_sheet("OP")
    _ecrire_entete(ws_op, ["N° OP", "Agent", "Montant total (Ar)", "Net payé (Ar)", "Statut"])
    for i, op in enumerate(data['ops'], start=2):
        ws_op.cell(row=i, column=1, value=op.numero_op)
        ws_op.cell(row=i, column=2, value=f"{op.agent_beneficiaire.matricule} {op.agent_beneficiaire.nom}")
        ws_op.cell(row=i, column=3, value=float(op.montant_total))
        ws_op.cell(row=i, column=4, value=float(op.montant_net_paye))
        ws_op.cell(row=i, column=5, value=op.get_statut_display())
        for col in range(1, 6):
            ws_op.cell(row=i, column=col).border = _BORDER
        ws_op.cell(row=i, column=3).number_format = '#,##0.00 "Ar"'
        ws_op.cell(row=i, column=4).number_format = '#,##0.00 "Ar"'
    _ajuster_colonnes(ws_op)

    # Feuille paiements
    ws_p = wb.create_sheet("Paiements")
    _ecrire_entete(ws_p, ["Date", "N° OP", "Montant (Ar)", "Mode", "Référence"])
    for i, p in enumerate(data['paiements'], start=2):
        ws_p.cell(row=i, column=1, value=p.date_paiement.replace(tzinfo=None))
        ws_p.cell(row=i, column=2, value=p.ordre_paiement.numero_op)
        ws_p.cell(row=i, column=3, value=float(p.montant))
        ws_p.cell(row=i, column=4, value=p.get_mode_paiement_display())
        ws_p.cell(row=i, column=5, value=p.reference or "—")
        ws_p.cell(row=i, column=1).number_format = 'DD/MM/YYYY HH:MM'
        ws_p.cell(row=i, column=3).number_format = '#,##0.00 "Ar"'
        for col in range(1, 6):
            ws_p.cell(row=i, column=col).border = _BORDER
    _ajuster_colonnes(ws_p)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ===================================================================
# EXPORTS CSV
# ===================================================================
# Format : séparateur ';' (Excel FR), UTF-8 avec BOM
# ===================================================================

import csv
from io import StringIO


def _csv_response_ready(buffer):
    """Encode le StringIO en UTF-8 avec BOM pour Excel."""
    return buffer.getvalue().encode('utf-8-sig')


def generer_etat_gardes_csv(date_debut=None, date_fin=None,
                            agent=None, mission=None, conforme=None):
    """Export CSV de l'état détaillé des gardes."""
    gardes = collecter_gardes_filtrees(
        date_debut, date_fin, agent, mission, conforme
    )

    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    writer.writerow([
        'Garde #', 'Agent', 'Matricule', 'Mission',
        'Date début', 'Date fin', 'Durée (min)', 'Partiels',
        'Conforme', 'Statut', 'Tarif partiel (Ar)',
    ])

    for g in gardes:
        writer.writerow([
            g.pk,
            f"{g.agent.nom} {g.agent.prenoms}",
            g.agent.matricule,
            g.mission.libelle if g.mission else '—',
            g.date_heure_debut.strftime('%d/%m/%Y %H:%M'),
            g.date_heure_fin.strftime('%d/%m/%Y %H:%M'),
            g.duree_minutes,
            g.nombre_partiels if g.nombre_partiels else '—',
            'Oui' if g.conforme else 'Non',
            g.get_statut_display(),
            g.tarif.montant_partiel if g.tarif else '—',
        ])

    return _csv_response_ready(buffer)


def generer_liste_op_csv(date_debut=None, date_fin=None, statut=None):
    """Export CSV de la liste des OP."""
    ops = collecter_ops_filtres(date_debut, date_fin, statut)

    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    writer.writerow([
        'N° OP', 'Agent', 'Matricule', 'Année',
        'Montant total (Ar)', 'Net payé (Ar)',
        'Statut', 'Clôturé', 'Créé le', 'Créé par',
    ])

    for op in ops:
        writer.writerow([
            op.numero_op,
            f"{op.agent_beneficiaire.nom} {op.agent_beneficiaire.prenoms}",
            op.agent_beneficiaire.matricule,
            op.annee,
            op.montant_total,
            op.montant_net_paye,
            op.get_statut_display(),
            'Oui' if op.cloture_administrative else 'Non',
            op.date_creation.strftime('%d/%m/%Y %H:%M'),
            op.id_utilisateur_createur.username,
        ])

    return _csv_response_ready(buffer)


def generer_rapport_periode_csv(date_debut, date_fin):
    """Export CSV du rapport sur période."""
    data = collecter_donnees_periode(date_debut, date_fin)

    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    # Section synthèse
    writer.writerow(['RAPPORT SUR PÉRIODE'])
    writer.writerow(['Du', date_debut.strftime('%d/%m/%Y')])
    writer.writerow(['Au', date_fin.strftime('%d/%m/%Y')])
    writer.writerow([])
    writer.writerow(['Indicateur', 'Valeur'])
    writer.writerow(['Nombre de gardes validées', data['nb_gardes']])
    writer.writerow(['Nombre de décomptes validés', data['nb_decomptes']])
    writer.writerow(['Nombre d\'OP créés', data['nb_ops']])
    writer.writerow(['Nombre de paiements enregistrés', data['nb_paiements']])
    writer.writerow(['Total décomptes (Ar)', data['total_decomptes']])
    writer.writerow(['Total OP (Ar)', data['total_ops']])
    writer.writerow(['Total paiements (Ar)', data['total_paiements']])

    # Section décomptes
    writer.writerow([])
    writer.writerow(['DÉCOMPTES VALIDÉS'])
    writer.writerow(['Agent', 'Période', 'Gardes', 'Montant (Ar)'])
    for d in data['decomptes']:
        writer.writerow([
            f"{d.agent.matricule} {d.agent.nom}",
            f"{d.mois:02d}/{d.annee}",
            d.nombre_gardes_retunes,
            d.montant_total,
        ])

    # Section OP
    writer.writerow([])
    writer.writerow(['ORDRES DE PAIEMENT'])
    writer.writerow(['N° OP', 'Agent', 'Montant total (Ar)', 'Net payé (Ar)', 'Statut'])
    for op in data['ops']:
        writer.writerow([
            op.numero_op,
            f"{op.agent_beneficiaire.matricule} {op.agent_beneficiaire.nom}",
            op.montant_total,
            op.montant_net_paye,
            op.get_statut_display(),
        ])

    # Section paiements
    writer.writerow([])
    writer.writerow(['PAIEMENTS'])
    writer.writerow(['Date', 'N° OP', 'Montant (Ar)', 'Mode', 'Référence'])
    for p in data['paiements']:
        writer.writerow([
            p.date_paiement.strftime('%d/%m/%Y %H:%M'),
            p.ordre_paiement.numero_op,
            p.montant,
            p.get_mode_paiement_display(),
            p.reference or '—',
        ])

    return _csv_response_ready(buffer)


def generer_historique_agent_csv(agent, date_debut=None, date_fin=None):
    """Export CSV de l'historique complet d'un agent."""
    data = collecter_donnees_historique_agent(agent, date_debut, date_fin)

    buffer = StringIO()
    writer = csv.writer(buffer, delimiter=';', quoting=csv.QUOTE_MINIMAL)

    # Informations agent
    writer.writerow(['FICHE AGENT'])
    writer.writerow(['Matricule', agent.matricule])
    writer.writerow(['Nom', agent.nom])
    writer.writerow(['Prénoms', agent.prenoms])
    writer.writerow(['Grade', agent.grade or '—'])
    writer.writerow(['Statut', agent.get_statut_agent_display()])
    writer.writerow(['Période analysée',
                     f"{data['debut'].strftime('%d/%m/%Y')} → {data['fin'].strftime('%d/%m/%Y')}"])

    # Gardes
    writer.writerow([])
    writer.writerow(['GARDES'])
    writer.writerow(['Garde #', 'Début', 'Fin', 'Durée (min)', 'Partiels',
                     'Conforme', 'Tarif (Ar)', 'Statut'])
    for g in data['gardes']:
        writer.writerow([
            g.pk,
            g.date_heure_debut.strftime('%d/%m/%Y %H:%M'),
            g.date_heure_fin.strftime('%d/%m/%Y %H:%M'),
            g.duree_minutes,
            g.nombre_partiels if g.nombre_partiels else '—',
            'Oui' if g.conforme else 'Non',
            g.tarif.montant_partiel if g.tarif else '—',
            g.get_statut_display(),
        ])

    # Décomptes
    writer.writerow([])
    writer.writerow(['DÉCOMPTES'])
    writer.writerow(['Décompte #', 'Période', 'Gardes', 'Non conformes',
                     'Montant (Ar)', 'Statut'])
    for d in data['decomptes']:
        writer.writerow([
            d.pk,
            f"{d.mois:02d}/{d.annee}",
            d.nombre_gardes_retunes,
            d.nombre_gardes_non_conformes_exclues,
            d.montant_total,
            d.get_statut_display(),
        ])

    # OP + Paiements
    writer.writerow([])
    writer.writerow(['ORDRES DE PAIEMENT'])
    writer.writerow(['N° OP', 'Date OP', 'Montant total (Ar)', 'Net payé (Ar)',
                     'Statut OP', 'Clôturé', 'Paiement date',
                     'Paiement montant (Ar)', 'Mode', 'Statut paiement'])
    for op in data['ordres']:
        paiements_op = list(op.paiements.all())
        if not paiements_op:
            writer.writerow([
                op.numero_op,
                op.date_creation.strftime('%d/%m/%Y'),
                op.montant_total,
                op.montant_net_paye,
                op.get_statut_display(),
                'Oui' if op.cloture_administrative else 'Non',
                '—', '—', '—', '—',
            ])
        else:
            for p in paiements_op:
                writer.writerow([
                    op.numero_op,
                    op.date_creation.strftime('%d/%m/%Y'),
                    op.montant_total,
                    op.montant_net_paye,
                    op.get_statut_display(),
                    'Oui' if op.cloture_administrative else 'Non',
                    p.date_paiement.strftime('%d/%m/%Y %H:%M'),
                    p.montant,
                    p.get_mode_paiement_display(),
                    p.get_statut_display(),
                ])

    return _csv_response_ready(buffer)


# ===================================================================
# QR CODE (Lot UX-1.2)
# ===================================================================

import base64
from io import BytesIO as _BytesIO


def _generer_qr_base64(data, taille_px=6):
    """
    Génère un QR code en PNG et le retourne encodé en base64.

    Paramètres :
    - data : texte à encoder
    - taille_px : taille de chaque module en pixels (défaut 6)

    Retourne : chaîne 'data:image/png;base64,...' prête à utiliser dans un
    template HTML (WeasyPrint accepte les images base64).
    """
    import segno

    qr = segno.make(data, error='M')
    buffer = _BytesIO()
    qr.save(buffer, kind='png', scale=taille_px, border=1)
    buffer.seek(0)
    encoded = base64.b64encode(buffer.getvalue()).decode('ascii')
    return f"data:image/png;base64,{encoded}"
