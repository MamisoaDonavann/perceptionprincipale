"""
Vues de l'application paiements.
"""

from copy import copy
from django.contrib import messages
from django.contrib.auth.decorators import login_required, permission_required
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Q
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST
from django.http import HttpResponse
from django.template.loader import render_to_string
from django.conf import settings
from django.utils import timezone  # 👈 IMPORT POUR LE FUSEAU HORAIRE
from weasyprint import HTML
from openpyxl import load_workbook
from io import BytesIO
import os
import math

from .forms import (
    ConfirmationReceptionForm, MotifForm, OrdrePaiementForm,
    PaiementForm, RetirerPeriodeForm,
)
from .models import ConfirmationReception, OrdrePaiement, Paiement
from .services import (
    annuler_confirmation, annuler_ordre_paiement, annuler_paiement,
    cloturer_administrativement, confirmer_reception,
    creer_ordre_paiement, enregistrer_paiement,
    retirer_periode_ordre, valider_ordre_paiement,
)
from decomptes.models import LigneDecompte

def _handle_validation_error(request, e):
    if hasattr(e, 'message_dict'):
        for field, errors in e.message_dict.items():
            for err in errors:
                messages.error(request, f"{field} : {err}")
    else:
        for msg in e.messages:
            messages.error(request, msg)

# -------------------------------------------------------------------
# Ordres de paiement
# -------------------------------------------------------------------

@login_required
@permission_required('paiements.consulter_ordre_paiement', raise_exception=True)
def op_list(request):
    qs = OrdrePaiement.objects.select_related('agent_beneficiaire', 'id_utilisateur_createur', 'id_utilisateur_validateur')
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(numero_op__icontains=q) | Q(agent_beneficiaire__matricule__icontains=q) | Q(agent_beneficiaire__nom__icontains=q))
    statut = request.GET.get('statut', '').strip()
    if statut in dict(OrdrePaiement.STATUTS_CHOICES):
        qs = qs.filter(statut=statut)
    annee = request.GET.get('annee', '').strip()
    if annee.isdigit():
        qs = qs.filter(annee=int(annee))
    qs = qs.order_by('-annee', '-numero_op')
    paginator = Paginator(qs, 20)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'paiements/list.html', {
        'page': page, 'q': q, 'statut': statut, 'annee': annee,
        'statuts_choices': OrdrePaiement.STATUTS_CHOICES, 'total': qs.count(),
    })

@login_required
@permission_required('paiements.consulter_ordre_paiement', raise_exception=True)
def op_detail(request, pk):
    op = get_object_or_404(OrdrePaiement.objects.select_related('agent_beneficiaire', 'id_utilisateur_createur', 'id_utilisateur_validateur', 'id_utilisateur_cloture'), pk=pk)
    periodes = op.periodes.select_related('decompte').order_by('decompte__annee', 'decompte__mois')
    paiements = op.paiements.select_related('id_utilisateur_enregistreur', 'id_utilisateur_annulateur').order_by('-date_paiement')
    paiements_actifs = paiements.filter(statut=Paiement.STATUT_ENREGISTRE)
    for p in paiements_actifs:
        p.confirmation_active = p.confirmations.filter(statut=ConfirmationReception.STATUT_CONFIRME).first()
    return render(request, 'paiements/detail.html', {
        'op': op, 'periodes': periodes, 'paiements': paiements, 'paiements_actifs': paiements_actifs,
        'peut_payer': op.statut == OrdrePaiement.STATUT_VALIDE,
        'peut_valider': op.statut == OrdrePaiement.STATUT_BROUILLON,
        'peut_annuler': op.statut in (OrdrePaiement.STATUT_BROUILLON, OrdrePaiement.STATUT_VALIDE, OrdrePaiement.STATUT_PAYE) and not op.paiements.filter(statut=Paiement.STATUT_ENREGISTRE).exists(),
        'peut_cloturer': (op.statut == OrdrePaiement.STATUT_PAYE and not op.cloture_administrative),
    })

@login_required
@permission_required('paiements.creer_ordre_paiement', raise_exception=True)
def op_create(request):
    if request.method == 'POST':
        form = OrdrePaiementForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data
            try:
                op, avertissements = creer_ordre_paiement(
                    agent=data['agent'], annee=data['annee'], numero_op=data['numero_op'],
                    decomptes=list(data['decomptes']), user=request.user,
                    observation=data.get('observation') or '', date_emission=data.get('date_emission') or None,
                    numero_be=data.get('numero_be') or '', titulaire_nom=data.get('titulaire_nom') or '',
                    titulaire_numero=data.get('titulaire_numero') or '', montant_lettres=data.get('montant_lettres') or '',
                    modalite_reglement=data.get('modalite_reglement') or '',
                )
                for a in avertissements: messages.warning(request, a)
                messages.success(request, f"Ordre de paiement {op.numero_op} créé — {op.montant_total} Ar.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = OrdrePaiementForm()
    return render(request, 'paiements/form.html', {'form': form, 'titre': "Nouvel ordre de paiement"})

@login_required
@permission_required('paiements.valider_ordre_paiement', raise_exception=True)
def op_valider(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    if request.method == 'POST':
        try:
            valider_ordre_paiement(op, request.user)
            messages.success(request, f"OP {op.numero_op} validé.")
            return redirect(reverse('paiements:detail', args=[op.pk]))
        except ValidationError as e:
            _handle_validation_error(request, e)
    return render(request, 'paiements/valider.html', {'op': op})

@login_required
@permission_required('paiements.creer_ordre_paiement', raise_exception=True)
def op_retirer_periode(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    if op.statut != OrdrePaiement.STATUT_BROUILLON:
        messages.error(request, "Seul un OP BROUILLON peut être modifié.")
        return redirect(reverse('paiements:detail', args=[op.pk]))
    if request.method == 'POST':
        form = RetirerPeriodeForm(request.POST, ordre=op)
        if form.is_valid():
            try:
                retirer_periode_ordre(op, form.cleaned_data['decompte'], request.user, form.cleaned_data['motif'])
                messages.success(request, "Période retirée.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = RetirerPeriodeForm(ordre=op)
    return render(request, 'paiements/retirer_periode.html', {'op': op, 'form': form})

@login_required
@permission_required('paiements.annuler_ordre_paiement', raise_exception=True)
def op_annuler(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    if request.method == 'POST':
        form = MotifForm(request.POST)
        if form.is_valid():
            try:
                annuler_ordre_paiement(op, request.user, form.cleaned_data['motif'])
                messages.success(request, f"OP {op.numero_op} annulé.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = MotifForm()
    return render(request, 'paiements/annuler.html', {'op': op, 'form': form})

@login_required
@permission_required('paiements.cloturer_ordre_paiement', raise_exception=True)
def op_cloturer(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    if request.method == 'POST':
        try:
            cloturer_administrativement(op, request.user)
            messages.success(request, f"OP {op.numero_op} clôturé administrativement.")
            return redirect(reverse('paiements:detail', args=[op.pk]))
        except ValidationError as e:
            _handle_validation_error(request, e)
    return render(request, 'paiements/cloturer.html', {'op': op})

@login_required
@permission_required('paiements.consulter_ordre_paiement', raise_exception=True)
def op_pdf_officiel(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    periodes = op.periodes.select_related('decompte').order_by('decompte__annee', 'decompte__mois')
    mois_noms = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
    mois_list = list(set([p.decompte.mois for p in periodes]))
    mois_list.sort()
    if len(mois_list) == 1: mois_str = f"{mois_noms[mois_list[0]]} {op.annee}"
    elif len(mois_list) > 1: mois_str = f"{mois_noms[mois_list[0]]} à {mois_noms[mois_list[-1]]} {op.annee}"
    else: mois_str = f"Année {op.annee}"
    context = {'op': op, 'mois_str': mois_str}
    html_string = render_to_string('rapports/pdf/ordre_paiement_officiel.html', context)
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="OP_Officiel_{op.numero_op}.pdf"'
    HTML(string=html_string, base_url=request.build_absolute_uri()).write_pdf(response)
    return response

@login_required
@permission_required('paiements.consulter_ordre_paiement', raise_exception=True)
def op_excel_officiel(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    
    periodes = op.periodes.select_related('decompte').order_by('decompte__annee', 'decompte__mois')
    mois_noms = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
    mois_list = list(set([p.decompte.mois for p in periodes]))
    mois_list.sort()
    if len(mois_list) == 1: mois_str = f"{mois_noms[mois_list[0]]} {op.annee}"
    elif len(mois_list) > 1: mois_str = f"{mois_noms[mois_list[0]]} à {mois_noms[mois_list[-1]]} {op.annee}"
    else: mois_str = f"Année {op.annee}"

    chemin_modele = os.path.join(settings.BASE_DIR, 'static', 'excel', 'modele_op.xlsx')
    
    if not os.path.exists(chemin_modele):
        messages.error(request, "Le fichier modèle Excel est introuvable dans static/excel/")
        return redirect(reverse('paiements:detail', args=[op.pk]))

    wb = load_workbook(chemin_modele)
    
    if 'SAISIE' in wb.sheetnames:
        ws = wb['SAISIE']
    else:
        ws = wb.active

    ws['B12'] = mois_str
    ws['B15'] = float(op.montant_total)
    ws['B16'] = op.montant_lettres or ""
    
    nom_complet = f"{op.agent_beneficiaire.nom} {op.agent_beneficiaire.prenoms}"
    ws['B19'] = op.titulaire_nom if op.titulaire_nom else nom_complet
    ws['B20'] = op.titulaire_numero if op.titulaire_numero else op.agent_beneficiaire.matricule
    ws['B21'] = op.agent_beneficiaire.cin or ""
    ws['B22'] = op.agent_beneficiaire.adresse or "CSP MAROVOAY"
    
    ws['B24'] = "Marovoay"
    ws['B25'] = op.annee
    
    date_ed = op.date_emission if op.date_emission else op.date_creation.date()
    ws['B26'] = date_ed.strftime('%d/%m/%Y')
    
    objet_texte = (
        f"Paiement d'indemnité(s) de mission intérieure de garde permanente\n"
        f"a/n de {nom_complet}\n"
        f"au titre du mois de (d') {mois_str}\n\n"
        f"auprès de la PP Marovoay"
    )
    ws['B28'] = objet_texte

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="OP_{op.numero_op.replace("/", "-")}.xlsx"'
    return response

# -------------------------------------------------------------------
# Génération Excel (Etats de décompte groupés - 2 par page)
# -------------------------------------------------------------------

def _trouver_cellule(ws, recherche):
    """Trouve une cellule contenant le texte (tolérant aux formules et symboles)."""
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and recherche in str(cell.value):
                return cell
    return None

def _remplacer_texte_global(ws, ancien_texte, nouveau_texte):
    """Remplace un texte dans TOUTES les cellules (tolérant aux sauts de ligne)."""
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                # On supprime les sauts de ligne pour la comparaison
                val_clean = cell.value.replace('\n', '').replace('\r', '')
                if ancien_texte in val_clean:
                    cell.value = val_clean.replace(ancien_texte, str(nouveau_texte) if nouveau_texte is not None else "")

def _remplir_tableau(ws, op, suffixe, balise_fin_nom):
    """Remplit un bloc de tableau en gérant l'insertion/suppression dynamique de lignes."""
    periodes = op.periodes.filter(reservation_active=True)
    
    mois_noms = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
    mois_list = list(set([p.decompte.mois for p in periodes]))
    mois_list.sort()
    if len(mois_list) == 1: mois_str = f"{mois_noms[mois_list[0]]}"
    elif len(mois_list) > 1: mois_str = f"{mois_noms[mois_list[0]]} à {mois_noms[mois_list[-1]]}"
    else: mois_str = "Indéfini"

    nom_complet = f"{op.agent_beneficiaire.nom} {op.agent_beneficiaire.prenoms}"

    # 1. Remplacer les textes globaux
    _remplacer_texte_global(ws, f"[MOIS{suffixe}]", mois_str)
    _remplacer_texte_global(ws, f"[ANNEE{suffixe}]", op.annee)
    _remplacer_texte_global(ws, f"[NOM{suffixe}]", nom_complet)
    _remplacer_texte_global(ws, f"[LETTRES{suffixe}]", f"{op.montant_lettres}")

    decompte_ids = [p.decompte_id for p in periodes]
    lignes = LigneDecompte.objects.filter(decompte_id__in=decompte_ids).select_related('garde').order_by('garde__date_heure_debut')

    cell_debut = _trouver_cellule(ws, f"[DEBUT{suffixe}]")
    cell_fin = _trouver_cellule(ws, f"[{balise_fin_nom}]")
    
    if not cell_debut or not cell_fin: 
        return 

    ligne_depart_data = cell_debut.row
    ligne_fin_max = cell_fin.row

    nb_lignes_necessaires = len(lignes)
    nb_lignes_disponibles = ligne_fin_max - ligne_depart_data

    # 2. Ajustement du nombre de lignes
    if nb_lignes_necessaires > nb_lignes_disponibles:
        diff = nb_lignes_necessaires - nb_lignes_disponibles
        ws.insert_rows(ligne_fin_max, amount=diff)
        for i in range(diff):
            target_row = ligne_fin_max + i
            for col in range(1, 8):
                source_cell = ws.cell(row=ligne_depart_data, column=col)
                target_cell = ws.cell(row=target_row, column=col)
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
            if ligne_depart_data in ws.row_dimensions:
                ws.row_dimensions[target_row].height = ws.row_dimensions[ligne_depart_data].height
                
    elif nb_lignes_disponibles > nb_lignes_necessaires:
        diff = nb_lignes_disponibles - nb_lignes_necessaires
        ws.delete_rows(ligne_depart_data + nb_lignes_necessaires, amount=diff)

    # 3. Remplissage des données
    total_partiels = 0
    current_row = ligne_depart_data
    
    for l in lignes:
        # 👇 CORRECTION DU FUSEAU HORAIRE (UTC vers Heure Locale)
        debut = l.garde.date_heure_debut
        fin = l.garde.date_heure_fin
        if timezone.is_aware(debut):
            debut = timezone.localtime(debut)
        if timezone.is_aware(fin):
            fin = timezone.localtime(fin)
            
        ws[f"A{current_row}"] = debut.strftime('%d/%m/%Y %H:%M')
        ws[f"B{current_row}"] = fin.strftime('%d/%m/%Y %H:%M')
        ws[f"C{current_row}"] = nom_complet
        ws[f"D{current_row}"] = op.agent_beneficiaire.matricule
        ws[f"E{current_row}"] = l.nombre_partiels_snapshot
        ws[f"F{current_row}"] = float(l.montant_partiel_snapshot)
        ws[f"G{current_row}"] = float(l.montant_ligne)

        total_partiels += l.nombre_partiels_snapshot
        current_row += 1

    # 4. Nettoyer les balises restantes (FIN_TABLEAU, etc.)
    _remplacer_texte_global(ws, f"[DEBUT{suffixe}]", "")
    _remplacer_texte_global(ws, f"[FIN{suffixe}]", "")
    _remplacer_texte_global(ws, f"[MATRICULE{suffixe}]", "")
    _remplacer_texte_global(ws, f"[{balise_fin_nom}]", "")

    # 5. Remplacer les totaux (avec sécurité pour les balises coupées)
    _remplacer_texte_global(ws, f"[TOTAL_GENERAL{suffixe}]", float(op.montant_total))
    
    # 👇 CORRECTION MAJEURE : On cherche [TOTAL_PARTIEL] même s'il y a des sauts de ligne
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str):
                clean_val = cell.value.replace('\n', '').replace('\r', '').replace(' ', '')
                if f"TOTAL_PARTIEL{suffixe}" in clean_val:
                    cell.value = int(total_partiels)
                    cell.number_format = 'General' # Force l'affichage du nombre


@login_required
@permission_required('paiements.consulter_ordre_paiement', raise_exception=True)
def op_excel_etat_decompte_bulk(request):
    """Génère le fichier Excel avec les états de décompte (2 par page)."""
    if request.method != 'POST':
        return redirect('paiements:list')

    op_ids = request.POST.getlist('op_ids')
    if not op_ids:
        messages.error(request, "Veuillez sélectionner au moins un OP.")
        return redirect('paiements:list')

    ops = list(OrdrePaiement.objects.filter(pk__in=op_ids).order_by('pk'))
    
    chemin_modele = os.path.join(settings.BASE_DIR, 'static', 'excel', 'modele_etat_decompte.xlsx')
    if not os.path.exists(chemin_modele):
        messages.error(request, "Le fichier modèle Excel est introuvable dans static/excel/")
        return redirect('paiements:list')

    nb_pages = math.ceil(len(ops) / 2)
    final_wb = None

    for i in range(nb_pages):
        page_wb = load_workbook(chemin_modele)
        ws = page_wb.active
        
        op1 = ops[i*2]
        op2 = ops[i*2 + 1] if i*2 + 1 < len(ops) else None

        # Remplissage du 1er tableau
        _remplir_tableau(ws, op1, "1", "FIN_TABLEAU1")
        
        # Remplissage du 2ème tableau
        if op2:
            _remplir_tableau(ws, op2, "2", "FIN_TABLEAU2")
        else:
            # Si pas de 2ème OP, on supprime COMPLÈTEMENT le 2ème tableau
            row_to_delete = None
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and "[MOIS2]" in cell.value:
                        row_to_delete = cell.row
                        break
                if row_to_delete: break
            
            if row_to_delete:
                start_row = row_to_delete
                for r in range(row_to_delete, 0, -1):
                    val = str(ws.cell(row=r, column=1).value) + str(ws.cell(row=r, column=2).value) + str(ws.cell(row=r, column=3).value)
                    if "ETAT DE DECOMPTE" in val:
                        start_row = r
                        break
                ws.delete_rows(start_row, ws.max_row - start_row + 1)

        ws.title = f"Page_{i+1}"
        
        if i == 0:
            final_wb = page_wb
        else:
            for sheet in page_wb.worksheets:
                final_wb._add_sheet(sheet)

    buffer = BytesIO()
    final_wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="Etats_Decomptes_Generes.xlsx"'
    return response

# -------------------------------------------------------------------
# Paiements
# -------------------------------------------------------------------

@login_required
@permission_required('paiements.enregistrer_paiement', raise_exception=True)
def paiement_create(request, pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    if request.method == 'POST':
        form = PaiementForm(request.POST, ordre=op)
        if form.is_valid():
            data = form.cleaned_data
            try:
                enregistrer_paiement(op, montant=data['montant'], mode=data['mode_paiement'], user=request.user, date_paiement=data['date_paiement'], reference=data.get('reference') or '')
                messages.success(request, f"Paiement de {data['montant']} Ar enregistré sur {op.numero_op}.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = PaiementForm(ordre=op)
    reste_a_payer = op.montant_total - op.montant_net_paye
    return render(request, 'paiements/paiement_form.html', {'op': op, 'form': form, 'reste_a_payer': reste_a_payer})

@login_required
@permission_required('paiements.annuler_paiement', raise_exception=True)
def paiement_annuler(request, pk, paiement_pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    paiement = get_object_or_404(Paiement, pk=paiement_pk, ordre_paiement=op)
    if request.method == 'POST':
        form = MotifForm(request.POST)
        if form.is_valid():
            try:
                annuler_paiement(paiement, request.user, form.cleaned_data['motif'])
                messages.success(request, "Paiement annulé.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = MotifForm()
    return render(request, 'paiements/paiement_annuler.html', {'op': op, 'paiement': paiement, 'form': form})

# -------------------------------------------------------------------
# Confirmations
# -------------------------------------------------------------------

@login_required
@permission_required('paiements.confirmer_reception', raise_exception=True)
def confirmation_create(request, pk, paiement_pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    paiement = get_object_or_404(Paiement, pk=paiement_pk, ordre_paiement=op)
    if request.method == 'POST':
        form = ConfirmationReceptionForm(request.POST, request.FILES)
        if form.is_valid():
            data = form.cleaned_data
            try:
                confirmer_reception(paiement, user=request.user, mode=data['mode'], justificatif=data['justificatif'], date_confirmation=data['date_confirmation'])
                messages.success(request, "Confirmation de réception enregistrée.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = ConfirmationReceptionForm()
    return render(request, 'paiements/confirmation_form.html', {'op': op, 'paiement': paiement, 'form': form})

@login_required
@permission_required('paiements.annuler_confirmation', raise_exception=True)
def confirmation_annuler(request, pk, paiement_pk, confirmation_pk):
    op = get_object_or_404(OrdrePaiement, pk=pk)
    paiement = get_object_or_404(Paiement, pk=paiement_pk, ordre_paiement=op)
    confirmation = get_object_or_404(ConfirmationReception, pk=confirmation_pk, paiement=paiement)
    if request.method == 'POST':
        form = MotifForm(request.POST)
        if form.is_valid():
            try:
                annuler_confirmation(confirmation, request.user, form.cleaned_data['motif'])
                messages.success(request, "Confirmation annulée.")
                return redirect(reverse('paiements:detail', args=[op.pk]))
            except ValidationError as e:
                _handle_validation_error(request, e)
    else:
        form = MotifForm()
    return render(request, 'paiements/confirmation_annuler.html', {'op': op, 'paiement': paiement, 'confirmation': confirmation, 'form': form})
